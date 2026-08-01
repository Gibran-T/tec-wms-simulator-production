#!/usr/bin/env python3
"""
TEC.WMS — M5 · Mission Superviseur d'exploitation
Présentation d'introduction visuelle (FR) — génération unique, hors modèles in-app.
"""

from __future__ import annotations

import os
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
# Paths & canvas
# ---------------------------------------------------------------------------

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "output" / "m5-pptx-assets"
OUT = ROOT / "output" / "TEC_WMS_M5_Mission_Superviseur_Exploitation.pptx"

# 16:9 widescreen
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Industrial palette — warehouse concrete / dock amber / night slate
CHARCOAL = RGBColor(0x1E, 0x2A, 0x38)
SLATE = RGBColor(0x2C, 0x3C, 0x4E)
STEEL = RGBColor(0x4A, 0x5D, 0x6E)
CONCRETE = RGBColor(0xEC, 0xE8, 0xE3)
PAPER = RGBColor(0xF8, 0xF6, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xD4, 0x8A, 0x2A)
SIGNAL = RGBColor(0x2F, 0x8F, 0x5B)
MUTED = RGBColor(0x6B, 0x75, 0x80)
SOFT = RGBColor(0x9A, 0xA3, 0xAB)


def img(name: str) -> str:
    path = ASSETS / name
    if not path.exists():
        raise FileNotFoundError(path)
    return str(path)


def set_run(run, text: str, size: int, color: RGBColor, bold: bool = False, font: str = "Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def add_text_box(slide, left, top, width, height, text, size=18, color=CHARCOAL,
                 bold=False, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, color, bold, font)
    return box


def add_multiline(slide, left, top, width, height, lines, size=15, color=CHARCOAL,
                  bold=False, align=PP_ALIGN.LEFT, spacing=1.15, font="Calibri"):
    """lines: list of str or (str, dict overrides)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        p.line_spacing = spacing
        if isinstance(line, tuple):
            text, opts = line
            run = p.add_run()
            set_run(
                run,
                text,
                opts.get("size", size),
                opts.get("color", color),
                opts.get("bold", bold),
                opts.get("font", font),
            )
        else:
            run = p.add_run()
            set_run(run, line, size, color, bold, font)
    return box


def fill_shape(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_shape(shape, color)
    return shape


def thin_bar(slide, left, top, width, color=AMBER, height=Pt(4)):
    return rect(slide, left, top, width, height, color)


def set_bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def picture_contain(slide, path: str, left, top, box_w, box_h, frame_color=None, pad=Inches(0.06)):
    """
    Place image fully visible inside a box (contain), never stretched.
    Optional thin professional frame around the photo area.
    """
    with Image.open(path) as im:
        iw, ih = im.size
    aspect = iw / ih
    box_aspect = box_w / box_h

    if aspect > box_aspect:
        # limited by width
        w = box_w - 2 * pad
        h = int(w / aspect)
    else:
        h = box_h - 2 * pad
        w = int(h * aspect)

    # center inside box
    x = left + (box_w - w) // 2
    y = top + (box_h - h) // 2

    if frame_color is not None:
        # outer frame slightly larger than image
        margin = Inches(0.04)
        frame = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            x - margin,
            y - margin,
            w + 2 * margin,
            h + 2 * margin,
        )
        fill_shape(frame, frame_color)

    pic = slide.shapes.add_picture(path, x, y, width=w, height=h)
    return pic


def footer(slide, page: int, total: int, light=False):
    c = SOFT if not light else RGBColor(0xA8, 0xB2, 0xBC)
    add_text_box(
        slide,
        Inches(0.5),
        Inches(7.15),
        Inches(10),
        Inches(0.28),
        "TEC.WMS — Gestion intégrée des stocks et de la performance logistique  ·  Module M5",
        size=10,
        color=c,
    )
    add_text_box(
        slide,
        Inches(11.2),
        Inches(7.15),
        Inches(1.6),
        Inches(0.28),
        f"{page} / {total}",
        size=10,
        color=c,
        align=PP_ALIGN.RIGHT,
    )


def chain_pills(slide, left, top, active=None, compact=False):
    """Render EXÉCUTER → MESURER → RÉCONCILIER → ARBITRER → DÉFENDRE."""
    steps = ["EXÉCUTER", "MESURER", "RÉCONCILIER", "ARBITRER", "DÉFENDRE"]
    pill_w = Inches(1.85) if not compact else Inches(1.7)
    pill_h = Inches(0.42) if not compact else Inches(0.38)
    gap = Inches(0.12)
    x = left
    for i, label in enumerate(steps):
        is_active = active is None or label == active or (isinstance(active, (list, set)) and label in active)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, pill_w, pill_h)
        fill_shape(shape, AMBER if is_active else STEEL if active else CHARCOAL)
        if not is_active and active is not None:
            fill_shape(shape, RGBColor(0x3A, 0x4A, 0x5C))
        tf = shape.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        set_run(run, label, 11 if not compact else 10, WHITE, True)
        # vertical center hack
        tf.paragraphs[0].space_before = Pt(6)
        x += pill_w + gap
        if i < len(steps) - 1:
            add_text_box(
                slide,
                x - gap + Inches(0.01),
                top + Inches(0.05),
                gap,
                Inches(0.32),
                "→",
                size=12,
                color=AMBER if active is None else SOFT,
                align=PP_ALIGN.CENTER,
                bold=True,
            )


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    TOTAL = 16

    # =====================================================================
    # 1 — Couverture
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, CHARCOAL)
    # Left panel text, right full photo contain
    rect(s, Inches(0), Inches(0), Inches(6.1), SLIDE_H, CHARCOAL)
    rect(s, Inches(6.1), Inches(0), Inches(7.233), SLIDE_H, RGBColor(0x14, 0x1C, 0x26))
    picture_contain(
        s,
        img("02_gls_truck_night.png"),
        Inches(6.25),
        Inches(0.35),
        Inches(6.7),
        Inches(6.8),
        frame_color=RGBColor(0x2A, 0x38, 0x48),
    )
    add_text_box(s, Inches(0.55), Inches(1.1), Inches(5.2), Inches(0.35),
                 "TEC.WMS", size=14, color=AMBER, bold=True)
    add_text_box(s, Inches(0.55), Inches(1.55), Inches(5.2), Inches(1.4),
                 "Module M5\nPrendre le quart", size=36, color=WHITE, bold=True)
    thin_bar(s, Inches(0.55), Inches(3.25), Inches(1.8), AMBER, Pt(5))
    add_multiline(
        s,
        Inches(0.55),
        Inches(3.55),
        Inches(5.1),
        Inches(2.2),
        [
            "Introduction visuelle à la mission",
            "de superviseur d'exploitation",
            "",
            "Simulation opérationnelle intégrée",
            "Gestion intégrée des stocks et",
            "de la performance logistique",
        ],
        size=15,
        color=SOFT,
    )
    add_text_box(s, Inches(0.55), Inches(6.55), Inches(5.1), Inches(0.4),
                 "Collège de la Concorde  ·  Peak Week", size=12, color=STEEL)
    footer(s, 1, TOTAL, light=True)

    # =====================================================================
    # 2 — Ce n'est plus un portefeuille
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, PAPER)
    thin_bar(s, Inches(0), Inches(0), SLIDE_W, AMBER, Pt(6))
    add_text_box(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.4),
                 "PASSAGE DE POSTURE", size=12, color=AMBER, bold=True)
    add_text_box(s, Inches(0.6), Inches(0.85), Inches(12), Inches(0.7),
                 "Vous n'observez plus un tableau.\nVous ouvrez un quart.", size=30, color=CHARCOAL, bold=True)

    # Two columns
    rect(s, Inches(0.6), Inches(2.1), Inches(5.7), Inches(4.4), CONCRETE)
    rect(s, Inches(6.9), Inches(2.1), Inches(5.8), Inches(4.4), CHARCOAL)
    add_text_box(s, Inches(0.9), Inches(2.35), Inches(5.1), Inches(0.4),
                 "Avant — lecture et classification", size=16, color=STEEL, bold=True)
    add_multiline(
        s,
        Inches(0.9),
        Inches(2.95),
        Inches(5.1),
        Inches(3.2),
        [
            "Analyser des indicateurs déjà constitués.",
            "Classer une situation.",
            "Proposer une orientation.",
            "",
            "Utile. Nécessaire.",
            "Mais ce n'est pas encore diriger le terrain.",
        ],
        size=16,
        color=CHARCOAL,
    )
    add_text_box(s, Inches(7.2), Inches(2.35), Inches(5.2), Inches(0.4),
                 "M5 — exécution puis décision", size=16, color=AMBER, bold=True)
    add_multiline(
        s,
        Inches(7.2),
        Inches(2.95),
        Inches(5.2),
        Inches(3.2),
        [
            "Faire avancer les opérations.",
            "Mesurer ce que votre session produit.",
            "Réconcilier les écarts avant de conclure.",
            "Arbitrer sous contrainte de temps.",
            "Défendre une décision avec des preuves.",
            "",
            "C'est la posture d'un superviseur de quart.",
        ],
        size=16,
        color=WHITE,
    )
    footer(s, 2, TOTAL)

    # =====================================================================
    # 3 — Le terrain réel
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, CHARCOAL)
    add_text_box(s, Inches(0.5), Inches(0.28), Inches(12), Inches(0.3),
                 "LE TERRAIN", size=12, color=AMBER, bold=True)
    add_text_box(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.45),
                 "Un centre de distribution n'attend pas. Il demande une lecture en continu.",
                 size=22, color=WHITE, bold=True)

    # Three vertical photos — full contain
    photos = [
        ("01_gls_freightline_dock.png", "Quai actif"),
        ("04_shunter_gls.png", "Cour et traction"),
        ("08_docks_143_144.png", "Portes et flux"),
    ]
    x0 = Inches(0.4)
    for i, (fn, caption) in enumerate(photos):
        left = x0 + i * Inches(4.25)
        picture_contain(s, img(fn), left, Inches(1.25), Inches(4.05), Inches(5.35), frame_color=SLATE)
        add_text_box(s, left, Inches(6.7), Inches(4.05), Inches(0.3),
                     caption, size=13, color=SOFT, align=PP_ALIGN.CENTER)
    footer(s, 3, TOTAL, light=True)

    # =====================================================================
    # 4 — Posture superviseur (pas une offre d'emploi)
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, PAPER)
    thin_bar(s, Inches(0), Inches(0), SLIDE_W, SIGNAL, Pt(6))
    add_text_box(s, Inches(0.55), Inches(0.35), Inches(7), Inches(0.3),
                 "POSTURE PROFESSIONNELLE", size=12, color=SIGNAL, bold=True)
    add_text_box(s, Inches(0.55), Inches(0.7), Inches(7.2), Inches(0.9),
                 "Superviseur d'exploitation :\nau cœur du quart", size=28, color=CHARCOAL, bold=True)
    add_multiline(
        s,
        Inches(0.55),
        Inches(1.85),
        Inches(6.6),
        Inches(4.8),
        [
            ("Ce matin, votre responsabilité n'est pas de « gérer en général ».", {"size": 15, "color": MUTED}),
            ("", {"size": 8}),
            ("Vous tenez le quai et les expéditions.", {"size": 16, "bold": True}),
            ("Vous coordonnez chauffeurs, manutentionnaires et priorités.", {"size": 16}),
            ("Vous contrôlez camions, remorques et équipements.", {"size": 16}),
            ("Vous suivez la ponctualité, la sécurité et les délais.", {"size": 16}),
            ("Vous traitez les écarts en temps réel.", {"size": 16}),
            ("Vous mesurez, vous décidez, vous documentez, vous transmettez.", {"size": 16}),
            ("", {"size": 10}),
            ("Référence terrain : profil Superviseur des opérations", {"size": 12, "color": STEEL}),
            ("(logistique & transport — quart matinal).", {"size": 12, "color": STEEL}),
        ],
        size=15,
        color=CHARCOAL,
    )
    picture_contain(
        s,
        img("14_dock_crew_pallet.png"),
        Inches(7.5),
        Inches(0.9),
        Inches(5.4),
        Inches(5.9),
        frame_color=CONCRETE,
    )
    footer(s, 4, TOTAL)

    # =====================================================================
    # 5 — La logique du quart (chaîne pédagogique)
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, PAPER)
    add_text_box(s, Inches(0.55), Inches(0.35), Inches(12), Inches(0.3),
                 "MÉTHODE DE TRAVAIL — MODULE M5", size=12, color=AMBER, bold=True)
    add_text_box(s, Inches(0.55), Inches(0.75), Inches(12), Inches(0.55),
                 "Une seule chaîne pour tenir le quart et réussir la mission",
                 size=26, color=CHARCOAL, bold=True)

    chain_pills(s, Inches(0.55), Inches(1.55), active=None)

    explanations = [
        ("EXÉCUTER", "Faire les opérations.\nAvancer le fret.\nTenir la séquence."),
        ("MESURER", "Lire le terrain.\nRelever les écarts.\nAncrer les preuves."),
        ("RÉCONCILIER", "Aligner vu / déclaré.\nCorriger avant de conclure.\nFermer la variance."),
        ("ARBITRER", "Choisir sous contrainte.\nProtéger le service.\nAssumer le compromis."),
        ("DÉFENDRE", "Justifier avec preuves.\nTransmettre le quart.\nRendre compte."),
    ]
    x = Inches(0.45)
    for title, body in explanations:
        rect(s, x, Inches(2.35), Inches(2.35), Inches(4.2), CONCRETE)
        thin_bar(s, x, Inches(2.35), Inches(2.35), AMBER, Pt(5))
        add_text_box(s, x + Inches(0.12), Inches(2.6), Inches(2.1), Inches(0.45),
                     title, size=13, color=CHARCOAL, bold=True, align=PP_ALIGN.CENTER)
        add_text_box(s, x + Inches(0.15), Inches(3.25), Inches(2.05), Inches(2.8),
                     body, size=13, color=STEEL, align=PP_ALIGN.CENTER)
        x += Inches(2.5)

    footer(s, 5, TOTAL)

    # =====================================================================
    # 6 — EXÉCUTER
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, CHARCOAL)
    chain_pills(s, Inches(0.4), Inches(0.25), active="EXÉCUTER", compact=True)
    add_text_box(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.5),
                 "EXÉCUTER — faire avancer le fret", size=28, color=WHITE, bold=True)
    add_text_box(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.45),
                 "Chargement, déchargement, mise à quai : le superviseur ne commente pas le flux — il le tient.",
                 size=15, color=SOFT)

    picture_contain(s, img("03_gls_trailer_open.png"), Inches(0.4), Inches(2.0), Inches(4.1), Inches(4.5), frame_color=SLATE)
    picture_contain(s, img("05_yard_tractor_dock117.png"), Inches(4.6), Inches(2.0), Inches(4.1), Inches(4.5), frame_color=SLATE)
    picture_contain(s, img("07_trailer_interior.png"), Inches(8.8), Inches(2.0), Inches(4.1), Inches(4.5), frame_color=SLATE)

    captions = [
        (0.4, "Contrôle du chargement"),
        (4.6, "Mise à quai et priorités"),
        (8.8, "Intégrité du fret"),
    ]
    for left, cap in captions:
        add_text_box(s, Inches(left), Inches(6.55), Inches(4.1), Inches(0.3),
                     cap, size=12, color=AMBER, align=PP_ALIGN.CENTER, bold=True)
    footer(s, 6, TOTAL, light=True)

    # =====================================================================
    # 7 — Quai comme poste de commandement
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, PAPER)
    add_text_box(s, Inches(0.5), Inches(0.3), Inches(6), Inches(0.3),
                 "SUPERVISION DU QUAI", size=12, color=AMBER, bold=True)
    add_text_box(s, Inches(0.5), Inches(0.65), Inches(6.2), Inches(0.7),
                 "Le quai est votre poste\nde commandement", size=26, color=CHARCOAL, bold=True)
    add_multiline(
        s,
        Inches(0.5),
        Inches(1.6),
        Inches(5.9),
        Inches(4.5),
        [
            "Affecter les portes.",
            "Lire les feux de quai et l'état des docks.",
            "Suivre les remorques en place et celles en attente.",
            "Anticiper les retards avant qu'ils n'explosent.",
            "",
            "Chaque porte ouverte ou fermée",
            "est une décision de capacité.",
            "",
            ("Dans M5, exécuter signifie enchaîner", {"size": 14, "color": STEEL}),
            ("les opérations sans rupture de séquence —", {"size": 14, "color": STEEL}),
            ("puis lire ce que cette exécution a produit.", {"size": 14, "color": STEEL}),
        ],
        size=16,
        color=CHARCOAL,
    )
    picture_contain(s, img("08_docks_143_144.png"), Inches(6.7), Inches(0.55), Inches(3.1), Inches(6.2), frame_color=CONCRETE)
    picture_contain(s, img("01_gls_freightline_dock.png"), Inches(10.0), Inches(0.55), Inches(3.0), Inches(6.2), frame_color=CONCRETE)
    footer(s, 7, TOTAL)

    # =====================================================================
    # 8 — Contrôle équipements
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, PAPER)
    thin_bar(s, Inches(0), Inches(0), SLIDE_W, STEEL, Pt(6))
    add_text_box(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.3),
                 "CONTRÔLE DES ÉQUIPEMENTS", size=12, color=STEEL, bold=True)
    add_text_box(s, Inches(0.5), Inches(0.65), Inches(12), Inches(0.5),
                 "Camions, remorques, tracteurs de cour — rien ne circule sans vigilance",
                 size=24, color=CHARCOAL, bold=True)

    picture_contain(s, img("04_shunter_gls.png"), Inches(0.4), Inches(1.45), Inches(4.0), Inches(5.2), frame_color=CONCRETE)
    picture_contain(s, img("02_gls_truck_night.png"), Inches(4.6), Inches(1.45), Inches(4.0), Inches(5.2), frame_color=CONCRETE)
    picture_contain(s, img("06_reefer_panel.png"), Inches(8.8), Inches(1.45), Inches(4.0), Inches(5.2), frame_color=CONCRETE)

    for left, label in [(0.4, "Tracteur de cour"), (4.6, "Ensemble routier"), (8.8, "Unité frigorifique")]:
        add_text_box(s, Inches(left), Inches(6.7), Inches(4.0), Inches(0.28),
                     label, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    footer(s, 8, TOTAL)

    # =====================================================================
    # 9 — MESURER
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, CHARCOAL)
    chain_pills(s, Inches(0.4), Inches(0.25), active="MESURER", compact=True)
    add_text_box(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.45),
                 "MESURER — lire ce que le terrain produit", size=26, color=WHITE, bold=True)

    picture_contain(s, img("06_reefer_panel.png"), Inches(0.35), Inches(1.45), Inches(3.7), Inches(5.2), frame_color=SLATE)
    picture_contain(s, img("14_dock_crew_pallet.png"), Inches(4.2), Inches(1.45), Inches(4.5), Inches(5.2), frame_color=SLATE)
    add_multiline(
        s,
        Inches(9.0),
        Inches(1.55),
        Inches(3.9),
        Inches(5.0),
        [
            ("Mesurer ≠ regarder un écran.", {"size": 15, "bold": True, "color": AMBER}),
            ("", {"size": 8}),
            ("C'est confronter une consigne à une lecture réelle : température, ponctualité, productivité, sécurité, écarts.", {"size": 13, "color": WHITE}),
            ("", {"size": 8}),
            ("Le panneau frigorifique et le suivi de quart font le même métier : ancrer des preuves avant toute conclusion.", {"size": 13, "color": SOFT}),
            ("", {"size": 10}),
            ("En M5, vos indicateurs viennent de votre propre session — pas d'un portefeuille.", {"size": 13, "bold": True, "color": WHITE}),
        ],
        size=13,
        color=WHITE,
    )
    footer(s, 9, TOTAL, light=True)

    # =====================================================================
    # 10 — RÉCONCILIER
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, PAPER)
    chain_pills(s, Inches(0.4), Inches(0.22), active="RÉCONCILIER", compact=True)
    add_text_box(s, Inches(0.5), Inches(0.8), Inches(12), Inches(0.45),
                 "RÉCONCILIER — aligner le vu et le déclaré", size=26, color=CHARCOAL, bold=True)
    add_text_box(s, Inches(0.5), Inches(1.3), Inches(12.3), Inches(0.4),
                 "Avant de décider, fermer l'écart. Une décision prise sur une variance ouverte est une décision fragile.",
                 size=14, color=MUTED)

    picture_contain(s, img("07_trailer_interior.png"), Inches(0.4), Inches(1.85), Inches(4.0), Inches(4.7), frame_color=CONCRETE)
    picture_contain(s, img("03_gls_trailer_open.png"), Inches(4.6), Inches(1.85), Inches(4.0), Inches(4.7), frame_color=CONCRETE)

    add_multiline(
        s,
        Inches(8.9),
        Inches(1.95),
        Inches(4.0),
        Inches(4.5),
        [
            ("Sur le terrain", {"size": 14, "bold": True, "color": AMBER}),
            ("Palettes, étiquettes, barres de maintien, espace libre, température, état de la remorque.", {"size": 13}),
            ("", {"size": 8}),
            ("Dans le système", {"size": 14, "bold": True, "color": AMBER}),
            ("Quantités, emplacements, écarts, statut d'étape, preuves du cycle.", {"size": 13}),
            ("", {"size": 8}),
            ("Votre métier", {"size": 14, "bold": True, "color": CHARCOAL}),
            ("Réconcilier d'abord. Décider ensuite. Justifier avec ce qui a été aligné.", {"size": 13}),
        ],
        size=13,
        color=CHARCOAL,
    )
    footer(s, 10, TOTAL)

    # =====================================================================
    # 11 — ARBITRER
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, CHARCOAL)
    chain_pills(s, Inches(0.4), Inches(0.25), active="ARBITRER", compact=True)
    add_text_box(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.45),
                 "ARBITRER — décider sous contrainte", size=26, color=WHITE, bold=True)

    picture_contain(s, img("05_yard_tractor_dock117.png"), Inches(0.35), Inches(1.5), Inches(4.15), Inches(5.15), frame_color=SLATE)
    picture_contain(s, img("12_supervisor_night_dock.png"), Inches(4.6), Inches(1.5), Inches(4.15), Inches(5.15), frame_color=SLATE)

    add_multiline(
        s,
        Inches(9.0),
        Inches(1.55),
        Inches(3.9),
        Inches(5.0),
        [
            ("Le terrain impose des contraintes :", {"size": 14, "bold": True, "color": AMBER}),
            ("", {"size": 6}),
            ("• météo et adhérence", {"size": 13, "color": WHITE}),
            ("• portes limitées", {"size": 13, "color": WHITE}),
            ("• délais clients", {"size": 13, "color": WHITE}),
            ("• sécurité non négociable", {"size": 13, "color": WHITE}),
            ("• capacité d'équipe", {"size": 13, "color": WHITE}),
            ("", {"size": 10}),
            ("Arbitrer, c'est choisir", {"size": 14, "bold": True, "color": WHITE}),
            ("ce qui passe en premier,", {"size": 14, "color": WHITE}),
            ("ce qui attend,", {"size": 14, "color": WHITE}),
            ("et ce qui est protégé.", {"size": 14, "color": WHITE}),
            ("", {"size": 10}),
            ("Sans liste magique.", {"size": 13, "color": SOFT}),
            ("Avec un raisonnement clair.", {"size": 13, "color": SOFT}),
        ],
        size=13,
        color=WHITE,
    )
    footer(s, 11, TOTAL, light=True)

    # =====================================================================
    # 12 — Sécurité
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, PAPER)
    thin_bar(s, Inches(0), Inches(0), SLIDE_W, SIGNAL, Pt(6))
    add_text_box(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.3),
                 "SÉCURITÉ ET CONFORMITÉ", size=12, color=SIGNAL, bold=True)
    add_text_box(s, Inches(0.5), Inches(0.65), Inches(12), Inches(0.5),
                 "Condition permanente — jamais une case à cocher en fin de quart",
                 size=24, color=CHARCOAL, bold=True)

    picture_contain(s, img("09_supervision_coaching.png"), Inches(0.4), Inches(1.4), Inches(6.2), Inches(5.2), frame_color=CONCRETE)
    picture_contain(s, img("12_supervisor_night_dock.png"), Inches(6.9), Inches(1.4), Inches(5.9), Inches(5.2), frame_color=CONCRETE)

    # overlay caption strip
    rect(s, Inches(0.4), Inches(6.2), Inches(12.4), Inches(0.55), CHARCOAL)
    add_text_box(
        s,
        Inches(0.6),
        Inches(6.3),
        Inches(12),
        Inches(0.4),
        "EPI · signaux de quai · communication claire · arrêt d'une manœuvre risquée  →  le service ne justifie jamais l'imprudence",
        size=13,
        color=WHITE,
    )
    footer(s, 12, TOTAL)

    # =====================================================================
    # 13 — Communication opérationnelle
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, PAPER)
    add_text_box(s, Inches(0.5), Inches(0.28), Inches(12), Inches(0.28),
                 "COMMUNICATION OPÉRATIONNELLE", size=12, color=AMBER, bold=True)
    add_text_box(s, Inches(0.5), Inches(0.6), Inches(12), Inches(0.5),
                 "Faire comprendre. Faire exécuter. Faire remonter.",
                 size=26, color=CHARCOAL, bold=True)

    # Three landscape scenes
    pics = [
        ("11_briefing_whiteboard.png", "Briefing d'équipe", "Aligner priorités et flux avant le rush."),
        ("09_supervision_coaching.png", "Coaching terrain", "Corriger en direct, près du fret."),
        ("13_supervisors_tablet.png", "Décision partagée", "Mesurer ensemble, arbitrer clairement."),
    ]
    x = Inches(0.35)
    for fn, title, blurb in pics:
        picture_contain(s, img(fn), x, Inches(1.3), Inches(4.15), Inches(3.7), frame_color=CONCRETE)
        add_text_box(s, x, Inches(5.15), Inches(4.15), Inches(0.35),
                     title, size=15, color=CHARCOAL, bold=True, align=PP_ALIGN.CENTER)
        add_text_box(s, x + Inches(0.1), Inches(5.55), Inches(3.95), Inches(0.7),
                     blurb, size=13, color=MUTED, align=PP_ALIGN.CENTER)
        x += Inches(4.3)
    footer(s, 13, TOTAL)

    # =====================================================================
    # 14 — DÉFENDRE / transmettre
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, CHARCOAL)
    chain_pills(s, Inches(0.4), Inches(0.25), active="DÉFENDRE", compact=True)
    add_text_box(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.45),
                 "DÉFENDRE — justifier, clôturer, transmettre", size=26, color=WHITE, bold=True)

    picture_contain(s, img("10_supervisors_handover.png"), Inches(0.35), Inches(1.5), Inches(6.3), Inches(5.1), frame_color=SLATE)
    add_multiline(
        s,
        Inches(7.0),
        Inches(1.6),
        Inches(5.8),
        Inches(5.0),
        [
            ("En fin de quart, vous ne « partez » pas.", {"size": 16, "bold": True, "color": AMBER}),
            ("Vous transmettez un état opérationnel.", {"size": 16, "color": WHITE}),
            ("", {"size": 10}),
            ("Ce que vous devez pouvoir défendre :", {"size": 14, "bold": True, "color": WHITE}),
            ("", {"size": 6}),
            ("• ce qui a été exécuté", {"size": 14, "color": SOFT}),
            ("• ce qui a été mesuré", {"size": 14, "color": SOFT}),
            ("• ce qui a été réconcilié", {"size": 14, "color": SOFT}),
            ("• ce qui a été arbitré", {"size": 14, "color": SOFT}),
            ("• ce qui reste ouvert pour le suivant", {"size": 14, "color": SOFT}),
            ("", {"size": 10}),
            ("M5 vous prépare à cette exigence :", {"size": 14, "color": WHITE}),
            ("décider avec des preuves de session,", {"size": 14, "color": WHITE}),
            ("en vos propres mots, sans phrase magique.", {"size": 14, "color": WHITE}),
        ],
        size=14,
        color=WHITE,
    )
    footer(s, 14, TOTAL, light=True)

    # =====================================================================
    # 15 — Contrat pédagogique M5 (sans réponses)
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, PAPER)
    thin_bar(s, Inches(0), Inches(0), SLIDE_W, AMBER, Pt(6))
    add_text_box(s, Inches(0.55), Inches(0.3), Inches(12), Inches(0.28),
                 "CONTRAT PÉDAGOGIQUE — MODULE M5", size=12, color=AMBER, bold=True)
    add_text_box(s, Inches(0.55), Inches(0.65), Inches(12), Inches(0.5),
                 "Ce que la mission vous demande — sans vous dicter la réponse",
                 size=24, color=CHARCOAL, bold=True)

    # Three competency cards
    cards = [
        ("C-M5-01", "Opérations intégrées", "Enchaîner le cycle de bout en bout.\nTenir la séquence.\nProduire des preuves propres."),
        ("C-M5-02", "Action corrective", "Traiter l'écart avant de conclure.\nRéconcilier, puis décider.\nNe pas décider sur une variance ouverte."),
        ("C-M5-03", "Décision capstone", "Lire l'instantané de votre session.\nArbitrer avec horizon.\nDéfendre un compromis explicite."),
    ]
    x = Inches(0.45)
    for code, title, body in cards:
        rect(s, x, Inches(1.4), Inches(4.0), Inches(3.5), CONCRETE)
        thin_bar(s, x, Inches(1.4), Inches(4.0), AMBER, Pt(5))
        add_text_box(s, x + Inches(0.25), Inches(1.65), Inches(3.5), Inches(0.3),
                     code, size=12, color=AMBER, bold=True)
        add_text_box(s, x + Inches(0.25), Inches(2.05), Inches(3.5), Inches(0.45),
                     title, size=18, color=CHARCOAL, bold=True)
        add_text_box(s, x + Inches(0.25), Inches(2.65), Inches(3.5), Inches(1.9),
                     body, size=14, color=STEEL)
        x += Inches(4.2)

    rect(s, Inches(0.45), Inches(5.15), Inches(12.4), Inches(1.45), CHARCOAL)
    add_multiline(
        s,
        Inches(0.75),
        Inches(5.35),
        Inches(11.8),
        Inches(1.15),
        [
            ("Règles du jeu", {"size": 14, "bold": True, "color": AMBER}),
            ("Utilisez vos propres mots. Ancrez-vous sur les preuves de votre session. Le système présente des faits — c'est vous qui produisez le diagnostic et la décision. Seuil de réussite : 70 / 100.", {"size": 14, "color": WHITE}),
        ],
        size=14,
        color=WHITE,
    )
    footer(s, 15, TOTAL)

    # =====================================================================
    # 16 — Ouverture de mission
    # =====================================================================
    s = prs.slides.add_slide(blank)
    set_bg(s, CHARCOAL)
    rect(s, Inches(0), Inches(0), Inches(6.5), SLIDE_H, RGBColor(0x12, 0x1A, 0x24))
    picture_contain(
        s,
        img("13_supervisors_tablet.png"),
        Inches(6.3),
        Inches(0.4),
        Inches(6.7),
        Inches(6.7),
        frame_color=SLATE,
    )
    add_text_box(s, Inches(0.55), Inches(1.5), Inches(5.5), Inches(0.35),
                 "OUVERTURE DE MISSION", size=13, color=AMBER, bold=True)
    add_text_box(s, Inches(0.55), Inches(2.0), Inches(5.5), Inches(1.2),
                 "Le quart commence.", size=36, color=WHITE, bold=True)
    thin_bar(s, Inches(0.55), Inches(3.35), Inches(1.6), AMBER, Pt(5))
    add_multiline(
        s,
        Inches(0.55),
        Inches(3.7),
        Inches(5.5),
        Inches(2.5),
        [
            "Exécutez. Mesurez. Réconciliez.",
            "Arbitrez. Défendez.",
            "",
            "Vos preuves viendront de votre session.",
            "Votre décision restera la vôtre.",
            "",
            "TEC.WMS · Module M5",
            "Simulation opérationnelle intégrée",
        ],
        size=15,
        color=SOFT,
    )
    footer(s, 16, TOTAL, light=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Saved: {OUT}")
    return str(OUT)


if __name__ == "__main__":
    build()
