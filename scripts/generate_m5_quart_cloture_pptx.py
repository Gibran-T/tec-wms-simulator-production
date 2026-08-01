#!/usr/bin/env python3
"""
TEC.WMS — M5 · Superviseur d'exploitation — Quart de clôture
Présentation d'introduction (FR) — architecture PRÉPARER → SUPERVISER → INTERVENIR → CLÔTURER
Nouvelle composition : ne reprend pas la structure du deck « Mission Superviseur ».
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "output" / "m5-closing-shift-assets"
OUT = ROOT / "output" / "TEC_WMS_M5_Superviseur_Quart_Cloture.pptx"
PREVIEW_DIR = ROOT / "output" / "m5-closing-shift-previews"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Palette distincte du deck précédent : béton clair + teal quai + anthracite
INK = RGBColor(0x1B, 0x24, 0x2E)
TEAL = RGBColor(0x1F, 0x6F, 0x6A)
TEAL_DEEP = RGBColor(0x14, 0x4A, 0x47)
SAND = RGBColor(0xF3, 0xF0, 0xEB)
MIST = RGBColor(0xE3, 0xE8, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COPPER = RGBColor(0xB8, 0x6B, 0x2D)
MUTED = RGBColor(0x5C, 0x68, 0x72)
SOFT = RGBColor(0x8A, 0x95, 0x9C)
LINE = RGBColor(0xC9, 0xD0, 0xCE)

ACTS = ["PRÉPARER", "SUPERVISER", "INTERVENIR", "CLÔTURER"]
TOTAL = 11


def img(name: str) -> str:
    path = ASSETS / name
    if not path.exists():
        raise FileNotFoundError(path)
    return str(path)


def set_run(run, text, size, color, bold=False, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def rect(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(sh, color)
    return sh


def textbox(slide, l, t, w, h, text, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    set_run(r, text, size, color, bold)
    return box


def multilines(slide, l, t, w, h, lines, size=15, color=INK, spacing=1.12):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(3)
        p.line_spacing = spacing
        if isinstance(line, tuple):
            txt, opts = line
            r = p.add_run()
            set_run(r, txt, opts.get("size", size), opts.get("color", color), opts.get("bold", False))
        else:
            r = p.add_run()
            set_run(r, line, size, color, False)
    return box


def picture_contain(slide, path, left, top, box_w, box_h, frame=None, pad=Inches(0.08)):
    """Full image visible — never cropped, never stretched."""
    with Image.open(path) as im:
        iw, ih = im.size
    aspect = iw / float(ih)
    box_aspect = box_w / float(box_h)
    if aspect > box_aspect:
        w = box_w - 2 * pad
        h = int(w / aspect)
    else:
        h = box_h - 2 * pad
        w = int(h * aspect)
    x = left + (box_w - w) // 2
    y = top + (box_h - h) // 2
    if frame is not None:
        m = Inches(0.035)
        fr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x - m, y - m, w + 2 * m, h + 2 * m)
        fill(fr, frame)
    return slide.shapes.add_picture(path, x, y, width=w, height=h)


def footer(slide, n, light=False):
    c = SOFT if not light else RGBColor(0x9A, 0xA8, 0xA6)
    textbox(
        slide,
        Inches(0.45),
        Inches(7.12),
        Inches(10.5),
        Inches(0.28),
        "TEC.WMS — Gestion intégrée des stocks et de la performance logistique  ·  Module M5",
        size=10,
        color=c,
    )
    textbox(slide, Inches(11.3), Inches(7.12), Inches(1.5), Inches(0.28), f"{n}/{TOTAL}", size=10, color=c, align=PP_ALIGN.RIGHT)


def act_ribbon(slide, left, top, active=None, width=Inches(12.4)):
    """Horizontal process — chevrons with readable inactive contrast."""
    n = len(ACTS)
    gap = Inches(0.1)
    unit = (width - gap * (n - 1)) / n
    x = left
    for act in ACTS:
        on = active is None or act == active
        shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, top, unit, Inches(0.48))
        fill(shape, TEAL if on else RGBColor(0xD5, 0xDC, 0xDA))
        tf = shape.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        # Inactive: darker ink on light fill for classroom projection
        set_run(r, act, 12, WHITE if on else INK, True)
        p.space_before = Pt(8)
        x += unit + gap


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ------------------------------------------------------------------ 1
    s = prs.slides.add_slide(blank)
    bg(s, INK)
    rect(s, Inches(0), Inches(0), Inches(0.18), SLIDE_H, TEAL)
    picture_contain(s, img("03_truck_night_gls.png"), Inches(7.1), Inches(0.35), Inches(5.8), Inches(6.7), frame=TEAL_DEEP)
    textbox(s, Inches(0.55), Inches(1.35), Inches(6.2), Inches(0.35), "TEC.WMS  ·  MODULE M5", size=13, color=TEAL, bold=True)
    textbox(s, Inches(0.55), Inches(1.85), Inches(6.2), Inches(1.6), "Superviseur d'exploitation\nQuart de clôture", size=34, color=WHITE, bold=True)
    rect(s, Inches(0.55), Inches(3.7), Inches(1.5), Pt(4), COPPER)
    multilines(
        s,
        Inches(0.55),
        Inches(4.05),
        Inches(6.0),
        Inches(2.0),
        [
            "Introduction à la mission opérationnelle",
            "de supervision — entreposage, transport",
            "et distribution.",
            "",
            "PRÉPARER  →  SUPERVISER  →  INTERVENIR  →  CLÔTURER",
        ],
        size=15,
        color=SOFT,
    )
    footer(s, 1, light=True)

    # ------------------------------------------------------------------ 2
    s = prs.slides.add_slide(blank)
    bg(s, SAND)
    textbox(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.3), "CONTEXTE DE QUART", size=12, color=TEAL, bold=True)
    textbox(s, Inches(0.5), Inches(0.7), Inches(12), Inches(0.55), "Le quart de clôture n'est pas une fin. C'est une responsabilité.", size=26, color=INK, bold=True)
    picture_contain(s, img("05_freightline_dock.png"), Inches(0.45), Inches(1.5), Inches(4.3), Inches(5.1), frame=MIST)
    picture_contain(s, img("08_docks_143_144.png"), Inches(4.95), Inches(1.5), Inches(4.0), Inches(5.1), frame=MIST)
    multilines(
        s,
        Inches(9.2),
        Inches(1.6),
        Inches(3.7),
        Inches(4.8),
        [
            ("Sur le terrain", {"bold": True, "size": 15, "color": TEAL}),
            "Portes actives ou libres.",
            "Remorques en place.",
            "Équipes encore en mouvement.",
            "Délais encore ouverts.",
            "",
            ("Votre rôle", {"bold": True, "size": 15, "color": TEAL}),
            "Tenir le service jusqu'à la transmission — sans laisser d'angle mort pour le prochain responsable.",
        ],
        size=14,
        color=INK,
    )
    footer(s, 2)

    # ------------------------------------------------------------------ 3
    s = prs.slides.add_slide(blank)
    bg(s, SAND)
    textbox(s, Inches(0.5), Inches(0.3), Inches(7), Inches(0.28), "POSTURE PROFESSIONNELLE", size=12, color=TEAL, bold=True)
    textbox(s, Inches(0.5), Inches(0.65), Inches(7), Inches(0.7), "Ce que tient un superviseur\nd'exploitation", size=28, color=INK, bold=True)
    multilines(
        s,
        Inches(0.5),
        Inches(1.7),
        Inches(6.5),
        Inches(4.8),
        [
            "• Quai, expéditions, affectations et priorités",
            "• Mouvement du fret — chargement et déchargement",
            "• Camions, remorques et équipements",
            "• Chauffeurs et équipes de soutien",
            "• Sécurité, conformité et heures de service",
            "• Écarts et incidents en temps réel",
            "• Suivi de performance et documentation",
            "• Transmission fiable en fin de quart",
            "",
            ("Sélection issue du profil Superviseur des opérations", {"size": 12, "color": MUTED}),
            ("(logistique et transport) — utiles pour comprendre le M5.", {"size": 12, "color": MUTED}),
        ],
        size=16,
        color=INK,
    )
    picture_contain(s, img("09_supervisor_oversight.png"), Inches(7.3), Inches(0.9), Inches(5.5), Inches(5.7), frame=MIST)
    footer(s, 3)

    # ------------------------------------------------------------------ 4
    s = prs.slides.add_slide(blank)
    bg(s, SAND)
    textbox(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.28), "MISSION TEC.WMS — MODULE M5", size=12, color=TEAL, bold=True)
    textbox(s, Inches(0.5), Inches(0.75), Inches(12), Inches(0.7), "Une seule responsabilité de quart.\nPas trois exercices séparés.", size=28, color=INK, bold=True)
    # Cumulative strip
    stages = [
        ("Étape 1", "Tenir le cycle\nnominal"),
        ("Étape 2", "Traiter l'écart\navant de conclure"),
        ("Étape 3", "Décider et\ntransmettre"),
    ]
    x = Inches(0.5)
    for i, (label, body) in enumerate(stages):
        rect(s, x, Inches(2.0), Inches(3.7), Inches(2.6), WHITE)
        rect(s, x, Inches(2.0), Inches(3.7), Pt(5), TEAL if i < 2 else COPPER)
        textbox(s, x + Inches(0.25), Inches(2.25), Inches(3.2), Inches(0.35), label, size=13, color=TEAL, bold=True)
        textbox(s, x + Inches(0.25), Inches(2.75), Inches(3.2), Inches(1.4), body, size=20, color=INK, bold=True)
        if i < 2:
            textbox(s, x + Inches(3.55), Inches(3.0), Inches(0.4), Inches(0.4), "→", size=22, color=COPPER, bold=True, align=PP_ALIGN.CENTER)
        x += Inches(4.05)
    multilines(
        s,
        Inches(0.5),
        Inches(5.0),
        Inches(12.3),
        Inches(1.5),
        [
            "Dans le M5, chaque étape s'appuie sur ce que vous avez déjà exécuté, observé et documenté.",
            "Vos preuves viennent de votre session. Le système présente des faits — vous produisez le raisonnement.",
            "Aucune phrase magique. Aucune conclusion automatique.",
        ],
        size=15,
        color=MUTED,
    )
    footer(s, 4)

    # ------------------------------------------------------------------ 5
    s = prs.slides.add_slide(blank)
    bg(s, INK)
    textbox(s, Inches(0.5), Inches(0.4), Inches(12), Inches(0.3), "ARCHITECTURE DU QUART", size=12, color=TEAL, bold=True)
    textbox(s, Inches(0.5), Inches(0.8), Inches(12), Inches(0.5), "Quatre actes. Une continuité.", size=30, color=WHITE, bold=True)
    act_ribbon(s, Inches(0.5), Inches(1.7), active=None, width=Inches(12.3))

    defs = [
        ("PRÉPARER", "Priorités, ressources,\naffectations, briefing."),
        ("SUPERVISER", "Suivre le quai, le fret,\nles équipements, le rythme."),
        ("INTERVENIR", "Corriger, sécuriser,\narbitrer sous contrainte."),
        ("CLÔTURER", "Documenter, transmettre,\nlaisser un état fiable."),
    ]
    x = Inches(0.5)
    for title, body in defs:
        rect(s, x, Inches(2.6), Inches(2.95), Inches(3.5), TEAL_DEEP)
        textbox(s, x + Inches(0.2), Inches(2.9), Inches(2.55), Inches(0.4), title, size=14, color=COPPER, bold=True)
        textbox(s, x + Inches(0.2), Inches(3.5), Inches(2.55), Inches(2.0), body, size=16, color=WHITE)
        x += Inches(3.15)
    footer(s, 5, light=True)

    # ------------------------------------------------------------------ 6 PRÉPARER
    s = prs.slides.add_slide(blank)
    bg(s, SAND)
    act_ribbon(s, Inches(0.45), Inches(0.25), active="PRÉPARER", width=Inches(12.4))
    textbox(s, Inches(0.5), Inches(0.95), Inches(6.5), Inches(0.5), "PRÉPARER — aligner avant le rush", size=26, color=INK, bold=True)
    multilines(
        s,
        Inches(0.5),
        Inches(1.65),
        Inches(5.8),
        Inches(4.8),
        [
            "Avant que le flux n'accélère :",
            "",
            "• clarifier les priorités du reste de quart ;",
            "• confirmer portes, itinéraires et affectations ;",
            "• rappeler les règles de sécurité ;",
            "• s'assurer que chacun sait quoi faire ensuite.",
            "",
            ("Un briefing court vaut mieux", {"bold": True, "color": TEAL}),
            ("qu'une improvisation coûteuse.", {"bold": True, "color": TEAL}),
        ],
        size=16,
        color=INK,
    )
    picture_contain(s, img("10_briefing_whiteboard.png"), Inches(6.6), Inches(1.2), Inches(6.2), Inches(5.4), frame=MIST)
    footer(s, 6)

    # ------------------------------------------------------------------ 7 SUPERVISER
    s = prs.slides.add_slide(blank)
    bg(s, SAND)
    act_ribbon(s, Inches(0.45), Inches(0.25), active="SUPERVISER", width=Inches(12.4))
    textbox(s, Inches(0.5), Inches(0.95), Inches(12), Inches(0.45), "SUPERVISER — lire le terrain en continu", size=26, color=INK, bold=True)
    multilines(
        s,
        Inches(0.5),
        Inches(1.55),
        Inches(4.3),
        Inches(5.0),
        [
            ("Votre regard reste sur l'exécution.", {"bold": True, "size": 16, "color": TEAL}),
            "",
            "Fret en cours de chargement.",
            "Contenu réel des remorques.",
            "État de la charge avant départ.",
            "",
            "Superviser, c'est confronter ce qui devrait se passer à ce qui se passe — porte par porte, chargement par chargement.",
            "",
            ("Sans cette lecture, aucune intervention", {"size": 14, "color": MUTED}),
            ("ni aucune clôture n'est fiable.", {"size": 14, "color": MUTED}),
        ],
        size=15,
        color=INK,
    )
    picture_contain(s, img("01_trailer_open_gls.png"), Inches(5.1), Inches(1.45), Inches(3.85), Inches(5.2), frame=MIST)
    picture_contain(s, img("07_trailer_interior.png"), Inches(9.15), Inches(1.45), Inches(3.7), Inches(5.2), frame=MIST)
    footer(s, 7)

    # ------------------------------------------------------------------ 8 SUPERVISER suite / contraintes
    s = prs.slides.add_slide(blank)
    bg(s, SAND)
    act_ribbon(s, Inches(0.45), Inches(0.25), active="SUPERVISER", width=Inches(12.4))
    textbox(s, Inches(0.5), Inches(0.95), Inches(12), Inches(0.45), "SUPERVISER — délais, équipements, conformité", size=24, color=INK, bold=True)
    picture_contain(s, img("04_yard_tractor_saq.png"), Inches(0.4), Inches(1.55), Inches(4.2), Inches(5.1), frame=MIST)
    picture_contain(s, img("06_reefer_panel.png"), Inches(4.8), Inches(1.55), Inches(4.0), Inches(5.1), frame=MIST)
    multilines(
        s,
        Inches(9.1),
        Inches(1.7),
        Inches(3.8),
        Inches(4.8),
        [
            ("Contraintes du quart", {"bold": True, "size": 15, "color": TEAL}),
            "Météo, portes, température, sécurité, ponctualité.",
            "",
            ("Ce que vous suivez", {"bold": True, "size": 15, "color": TEAL}),
            "Tracteurs de cour, remorques, unités frigorifiques.",
            "Le rythme du quai.",
            "Les signaux de conformité.",
            "",
            "Superviser, c'est détecter tôt — avant qu'un écart ne devienne un incident.",
        ],
        size=14,
        color=INK,
    )
    footer(s, 8)

    # ------------------------------------------------------------------ 9 INTERVENIR
    s = prs.slides.add_slide(blank)
    bg(s, SAND)
    act_ribbon(s, Inches(0.45), Inches(0.25), active="INTERVENIR", width=Inches(12.4))
    textbox(s, Inches(0.5), Inches(0.95), Inches(12), Inches(0.4), "INTERVENIR — agir quand la situation l'exige", size=24, color=INK, bold=True)
    textbox(
        s,
        Inches(0.5),
        Inches(1.4),
        Inches(12.3),
        Inches(0.35),
        "Stopper une manœuvre  ·  recentrer une équipe  ·  protéger le service et la sécurité  ·  documenter l'action",
        size=13,
        color=TEAL,
        bold=True,
    )
    picture_contain(s, img("13_coaching_floor.png"), Inches(0.4), Inches(1.9), Inches(6.1), Inches(4.7), frame=MIST)
    picture_contain(s, img("11_dock_crew_night.png"), Inches(6.7), Inches(1.9), Inches(6.1), Inches(4.7), frame=MIST)
    footer(s, 9)

    # ------------------------------------------------------------------ 10 CLÔTURER
    s = prs.slides.add_slide(blank)
    bg(s, SAND)
    act_ribbon(s, Inches(0.45), Inches(0.25), active="CLÔTURER", width=Inches(12.4))
    textbox(s, Inches(0.5), Inches(0.95), Inches(12), Inches(0.45), "CLÔTURER — transmettre un état fiable", size=26, color=INK, bold=True)
    picture_contain(s, img("12_handover_tablet.png"), Inches(0.4), Inches(1.55), Inches(7.2), Inches(5.0), frame=MIST)
    multilines(
        s,
        Inches(7.9),
        Inches(1.7),
        Inches(4.9),
        Inches(4.8),
        [
            ("En fin de quart, vous laissez :", {"bold": True, "size": 15, "color": TEAL}),
            "",
            "• ce qui a été exécuté ;",
            "• ce qui a été mesuré ;",
            "• les écarts traités ou encore ouverts ;",
            "• les actions déjà engagées ;",
            "• ce que le prochain doit reprendre.",
            "",
            ("Clôturer, ce n'est pas partir.", {"bold": True, "color": INK}),
            ("C'est rendre le quart défendable.", {"bold": True, "color": INK}),
        ],
        size=15,
        color=INK,
    )
    footer(s, 10)

    # ------------------------------------------------------------------ 11 Ouverture M5
    s = prs.slides.add_slide(blank)
    bg(s, INK)
    rect(s, Inches(0), Inches(0), Inches(0.18), SLIDE_H, COPPER)
    textbox(s, Inches(0.55), Inches(1.2), Inches(12), Inches(0.3), "OUVERTURE DE MISSION", size=13, color=TEAL, bold=True)
    textbox(s, Inches(0.55), Inches(1.7), Inches(12), Inches(0.7), "Vous prenez le quart de clôture.", size=32, color=WHITE, bold=True)
    act_ribbon(s, Inches(0.55), Inches(2.7), active=None, width=Inches(12.2))
    multilines(
        s,
        Inches(0.55),
        Inches(3.6),
        Inches(12.2),
        Inches(2.6),
        [
            "Préparez. Supervisez. Intervenez. Clôturez.",
            "",
            "Dans TEC.WMS, le Module M5 vous place dans cette continuité :",
            "exécuter, suivre, traiter les écarts, puis justifier une décision",
            "à partir des preuves de votre propre session.",
            "",
            "Vos mots. Votre raisonnement. Votre responsabilité.",
        ],
        size=16,
        color=SOFT,
    )
    footer(s, 11, light=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Saved: {OUT}")
    return str(OUT)


if __name__ == "__main__":
    build()
