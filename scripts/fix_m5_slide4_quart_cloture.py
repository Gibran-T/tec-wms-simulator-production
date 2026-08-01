#!/usr/bin/env python3
"""Correction urgente — slide 4 uniquement (4 actes officiels)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ORIG = Path(r"c:\Projetos\tec-wms-simulator-production\output\TEC_WMS_M5_Superviseur_Quart_Cloture.pptx")
OUT = Path(r"c:\Projetos\tec-wms-simulator-production\output\TEC_WMS_M5_Superviseur_Quart_Cloture_SLIDE4_CORRIGE.pptx")

INK = RGBColor(0x1B, 0x24, 0x2E)
TEAL = RGBColor(0x1F, 0x6F, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COPPER = RGBColor(0xB8, 0x6B, 0x2D)
MUTED = RGBColor(0x5C, 0x68, 0x72)

ACTS = [
    ("PRÉPARER", "Aligner le quart,\nles priorités et\nles ressources."),
    ("SUPERVISER", "Suivre la réalité\nde l'exploitation\net l'exécution."),
    ("INTERVENIR", "Traiter les écarts\net protéger la\ncontinuité opérationnelle."),
    ("CLÔTURER", "Consolider l'état\ndu quart et transmettre\nune information fiable."),
]


def set_run(run, text, size, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"


def fill_solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def set_textbox_text(shape, lines, size, color, bold=False, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.clear()
    p0.alignment = align
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = True
    for line in lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        r = p.add_run()
        set_run(r, line, size, color, bold)


def main():
    if not ORIG.exists():
        raise SystemExit("ORIGINAL_MISSING: " + str(ORIG))

    prs = Presentation(str(ORIG))
    if len(prs.slides) != 11:
        raise SystemExit(f"Unexpected slide count: {len(prs.slides)}")

    slide = prs.slides[3]
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    if not any("Étape 1" in t for t in texts):
        raise SystemExit("Unexpected slide 4 — Étape 1 missing; refusing to modify.")
    if not any("Étape 3" in t for t in texts):
        raise SystemExit("Unexpected slide 4 — Étape 3 missing; refusing to modify.")
    print("Verified original Étape labels present.")

    to_remove = []
    for sh in slide.shapes:
        keep = False
        if sh.has_text_frame:
            t = sh.text_frame.text
            if t.startswith("MISSION TEC.WMS"):
                keep = True
            elif t.startswith("Une seule responsabilité"):
                keep = True
            elif t.startswith("Dans le M5"):
                keep = True
            elif t.startswith("TEC.WMS — Gestion"):
                keep = True
            elif t.strip() == "4/11":
                keep = True
        if not keep:
            to_remove.append(sh)

    print(f"Removing {len(to_remove)} card-cluster shapes; keeping headers/footer/pedagogy.")
    for sh in to_remove:
        el = sh._element
        el.getparent().remove(el)

    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("Dans le M5"):
            set_textbox_text(
                sh,
                [
                    "Dans le M5, chaque acte s'appuie sur ce que vous avez déjà préparé, supervisé et documenté.",
                    "Vos preuves viennent de votre session. Le système présente des faits — vous produisez le raisonnement.",
                    "Aucune phrase magique. Aucune conclusion automatique.",
                ],
                size=15,
                color=MUTED,
                bold=False,
            )

    left0 = Inches(0.45)
    top = Emu(1828800)
    card_h = Emu(2377440)
    bar_h = Emu(63500)
    usable = Inches(12.4)
    gap = Inches(0.18)
    n = 4
    card_w = int((usable - gap * (n - 1)) / n)

    x = int(left0)
    for i, (title, body) in enumerate(ACTS):
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top, card_w, card_h)
        fill_solid(card, WHITE)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top, card_w, bar_h)
        fill_solid(bar, COPPER if i == n - 1 else TEAL)
        tb = slide.shapes.add_textbox(
            x + Inches(0.16), Emu(2057400), card_w - Inches(0.28), Inches(0.4)
        )
        set_textbox_text(tb, [title], size=13, color=TEAL, bold=True)
        bb = slide.shapes.add_textbox(
            x + Inches(0.16), Emu(2450000), card_w - Inches(0.28), Inches(1.7)
        )
        set_textbox_text(bb, body.split("\n"), size=15, color=INK, bold=True)
        x += card_w + int(gap)

    prs.save(str(OUT))
    print(f"Saved: {OUT}")

    prs2 = Presentation(str(OUT))
    s4 = prs2.slides[3]
    all_txt = "\n".join(sh.text_frame.text for sh in s4.shapes if sh.has_text_frame)
    print("--- SLIDE 4 TEXT ---")
    print(all_txt)
    assert "Étape 1" not in all_txt
    assert "Étape 2" not in all_txt
    assert "Étape 3" not in all_txt
    for act in ("PRÉPARER", "SUPERVISER", "INTERVENIR", "CLÔTURER"):
        assert act in all_txt, act
    assert "Pas trois exercices séparés" in all_txt

    prs_o = Presentation(str(ORIG))
    for i in range(11):
        if i == 3:
            continue
        if len(prs_o.slides[i].shapes) != len(prs2.slides[i].shapes):
            raise SystemExit(f"Slide {i+1} altered")
    print("OK: four acts present; other slides unchanged.")


if __name__ == "__main__":
    main()
